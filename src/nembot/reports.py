"""Report and notebook export for chat results.

The model writes the report as markdown and places query/plot results on
their own lines as ``[[R3]]`` (result ids are assigned when tools run). Every
export re-runs the referenced SQL, so files reflect the current data.
Formats: html, pdf, docx, pptx, xlsx, jupyter (.ipynb), marimo (.py).
"""
from dataclasses import dataclass
from datetime import datetime
import base64
import html
import io
import json
from pathlib import Path
import re

import altair
import markdown as markdown_lib
import vl_convert

from .attribution import CREDIT
from .chat import Result, _jsonable, _records, run_query

FORMATS = {'html': '.html', 'pdf': '.pdf', 'docx': '.docx', 'pptx': '.pptx', 'xlsx': '.xlsx',
           'jupyter': '.ipynb', 'marimo': '.py'}
NOTEBOOKS = ('jupyter', 'marimo')
TABLE_ROWS = 20  # rows of each result shown in documents (the xlsx and notebooks have everything)
MARKER = re.compile(r'^\s*\[\[(R\d+)\]\]\s*$')


@dataclass
class Item:
    """A result included in a report, re-run at export time."""
    id: str
    sql: str
    plot_spec: dict | None
    result: Result | None = None
    error: str | None = None
    spec: dict | None = None   # full Vega-Lite spec with data
    png: bytes | None = None


def slug(text: str) -> str:
    return re.sub(r'[^a-z0-9]+', '-', text.lower()).strip('-')[:60] or 'report'


def blocks(content: str) -> list[tuple[str, str]]:
    """Split markdown into ('text', md) and ('result', 'R3') blocks."""
    out, buffer = [], []
    for line in content.splitlines():
        if match := MARKER.match(line):
            if ''.join(buffer).strip():
                out.append(('text', '\n'.join(buffer).strip()))
            buffer = []
            out.append(('result', match[1]))
        else:
            buffer.append(line)
    if ''.join(buffer).strip():
        out.append(('text', '\n'.join(buffer).strip()))
    return out


def prepare(db, refs: list[str], results: dict, *, images: bool) -> dict[str, Item]:
    items = {}
    for ref in refs:
        if ref in items or ref not in results:
            continue
        stored = results[ref]
        item = Item(ref, stored['sql'], stored.get('plot_spec'))
        result = run_query(db, item.sql)
        if isinstance(result, str):
            item.error = result
        else:
            item.result = result
            if item.plot_spec:
                try:
                    item.spec = altair.Chart.from_dict({**item.plot_spec, 'data': {'values': _records(result)}},
                                                       validate=True).to_dict()
                    if images:
                        item.png = vl_convert.vegalite_to_png(item.spec, scale=2)
                except Exception as exc:
                    item.error = f'Plot failed: {exc}'
        items[ref] = item
    return items


def _cell(value) -> str:
    value = _jsonable(value)
    if value is None:
        return ''
    if isinstance(value, float):
        return f'{value:,.4f}'.rstrip('0').rstrip('.')
    return str(value)


# ---- HTML / PDF ----

_CSS = """
body { font: 15px/1.55 system-ui, sans-serif; color: #1f2328; max-width: 920px; margin: 32px auto; padding: 0 24px; }
h1 { font-size: 26px; margin-bottom: 4px; } .meta { color: #656d76; font-size: 13px; margin-bottom: 24px; }
table { border-collapse: collapse; font-size: 12px; margin: 8px 0; } th, td { border-bottom: 1px solid #d0d7de; padding: 3px 8px; text-align: left; }
td.num { text-align: right; } .note { color: #656d76; font-size: 12px; } .error { color: #cf222e; }
pre { background: #f6f8fa; padding: 8px; font-size: 11px; white-space: pre-wrap; } figure { margin: 12px 0; }
img { max-width: 100%; } .result { margin: 16px 0; page-break-inside: avoid; }
"""


def _html_table(result: Result) -> str:
    head = ''.join(f'<th>{html.escape(c)}</th>' for c in result.columns)
    rows = ''.join('<tr>' + ''.join(
        f'<td class="num">{html.escape(_cell(v))}</td>' if isinstance(v, (int, float)) and not isinstance(v, bool)
        else f'<td>{html.escape(_cell(v))}</td>' for v in row) + '</tr>' for row in result.rows[:TABLE_ROWS])
    note = (f'<p class="note">First {TABLE_ROWS} of {len(result.rows):,} rows.</p>'
            if len(result.rows) > TABLE_ROWS else '')
    return f'<table><tr>{head}</tr>{rows}</table>{note}'


def build_html(title: str, content: str, items: dict[str, Item], *, static: bool) -> str:
    parts, appendix, embeds = [], [], []
    for kind, value in blocks(content):
        if kind == 'text':
            parts.append(markdown_lib.markdown(value, extensions=['tables']))
            continue
        item = items.get(value)
        if item is None:
            parts.append(f'<p class="error">Unknown result {html.escape(value)}.</p>')
            continue
        body = f'<p class="error">{html.escape(item.error)}</p>' if item.error else ''
        if item.spec and static and item.png:
            body += f'<figure><img alt="{item.id}" src="data:image/png;base64,{base64.b64encode(item.png).decode()}"></figure>'
        elif item.spec:
            embeds.append((f'vis-{item.id}', item.spec))
            body += f'<figure id="vis-{item.id}"></figure>'
        elif item.result is not None:
            body += _html_table(item.result)
        parts.append(f'<div class="result">{body}<p class="note">Result {item.id}; SQL in the appendix.</p></div>')
        appendix.append(f'<h3>{item.id}</h3><pre>{html.escape(item.sql)}</pre>')
    scripts = ''
    if embeds:
        scripts = ('<script src="https://cdn.jsdelivr.net/npm/vega@6"></script>'
                   '<script src="https://cdn.jsdelivr.net/npm/vega-lite@6"></script>'
                   '<script src="https://cdn.jsdelivr.net/npm/vega-embed@7"></script><script>'
                   + ''.join(f'vegaEmbed("#{target}", {json.dumps(spec, default=str)}, {{actions: false}});'
                             for target, spec in embeds) + '</script>')
    meta = f'Generated by nembot on {datetime.now():%d %B %Y %H:%M} from AEMO MMS data.'
    return (f'<!doctype html><html><head><meta charset="utf-8"><title>{html.escape(title)}</title>'
            f'<style>{_CSS}</style></head><body><h1>{html.escape(title)}</h1><p class="meta">{meta}</p>'
            + ''.join(parts) + (f'<h2>Appendix: SQL</h2>{"".join(appendix)}' if appendix else '')
            + f'<p class="note credit">{html.escape(CREDIT)}</p>' + scripts + '</body></html>')


def build_pdf(title, content, items) -> bytes:
    import weasyprint
    return weasyprint.HTML(string=build_html(title, content, items, static=True)).write_pdf()


# ---- DOCX ----

_INLINE = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)')


def _runs(paragraph, text: str):
    for piece in _INLINE.split(text):
        if not piece:
            continue
        if piece.startswith('**'):
            paragraph.add_run(piece[2:-2]).bold = True
        elif piece.startswith('`'):
            paragraph.add_run(piece[1:-1]).font.name = 'Consolas'
        elif piece.startswith('*'):
            paragraph.add_run(piece[1:-1]).italic = True
        else:
            paragraph.add_run(piece)


def _docx_markdown(document, text: str):
    paragraph_lines = []

    def flush():
        if paragraph_lines:
            _runs(document.add_paragraph(), ' '.join(paragraph_lines))
            paragraph_lines.clear()
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            flush()
        elif match := re.match(r'(#{1,4})\s+(.*)', stripped):
            flush()
            document.add_heading(match[2], level=min(len(match[1]), 3))
        elif match := re.match(r'[-*]\s+(.*)', stripped):
            flush()
            _runs(document.add_paragraph(style='List Bullet'), match[1])
        elif match := re.match(r'\d+[.)]\s+(.*)', stripped):
            flush()
            _runs(document.add_paragraph(style='List Number'), match[1])
        else:
            paragraph_lines.append(stripped)
    flush()


def build_docx(title, content, items) -> bytes:
    import docx
    from docx.shared import Inches, Pt
    document = docx.Document()
    document.add_heading(title, level=0)
    document.add_paragraph(f'Generated by nembot on {datetime.now():%d %B %Y} from AEMO MMS data.').italic = True
    for kind, value in blocks(content):
        if kind == 'text':
            _docx_markdown(document, value)
            continue
        item = items.get(value)
        if item is None or item.error:
            document.add_paragraph(item.error if item else f'Unknown result {value}.')
            continue
        if item.png:
            document.add_picture(io.BytesIO(item.png), width=Inches(6.3))
        elif item.result is not None:
            rows = item.result.rows[:TABLE_ROWS]
            table = document.add_table(rows=len(rows) + 1, cols=len(item.result.columns))
            table.style = 'Light Grid Accent 1'
            for i, name in enumerate(item.result.columns):
                table.cell(0, i).text = name
            for r, row in enumerate(rows, 1):
                for i, value_ in enumerate(row):
                    table.cell(r, i).text = _cell(value_)
            if len(item.result.rows) > TABLE_ROWS:
                document.add_paragraph(f'First {TABLE_ROWS} of {len(item.result.rows):,} rows.').runs[0].font.size = Pt(8)
    document.add_heading('Appendix: SQL', level=1)
    for item in items.values():
        document.add_heading(item.id, level=3)
        for line in item.sql.splitlines():
            run = document.add_paragraph().add_run(line)
            run.font.name, run.font.size = 'Consolas', Pt(8)
    credit = document.add_paragraph().add_run(CREDIT)
    credit.italic, credit.font.size = True, Pt(8)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ---- PPTX: a title slide, then one slide per section (heading, findings, its first plot/table) ----

def _plain(text: str) -> str:
    return re.sub(r'[*`#]', '', text).strip()


def _sections(content: str) -> list[dict]:
    sections, current = [], {'title': None, 'lines': [], 'refs': []}
    for kind, value in blocks(content):
        if kind == 'result':
            current['refs'].append(value)
            continue
        for line in value.splitlines():
            if match := re.match(r'#{1,3}\s+(.*)', line.strip()):
                if current['title'] or current['lines'] or current['refs']:
                    sections.append(current)
                current = {'title': match[1], 'lines': [], 'refs': []}
            elif line.strip():
                current['lines'].append(_plain(re.sub(r'^\s*([-*]|\d+[.)])\s+', '', line)))
    if current['title'] or current['lines'] or current['refs']:
        sections.append(current)
    return sections


def build_pptx(title, content, items) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    def place(shape, top, height):  # the default template's placeholders are laid out for 4:3
        shape.left, shape.top, shape.width, shape.height = Inches(0.6), top, Inches(12.1), height

    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = title
    place(slide.shapes.title, Inches(2.2), Inches(1.8))
    slide.placeholders[1].text = f'nembot · Data: AEMO (NEMWEB) · {datetime.now():%d %B %Y}'
    place(slide.placeholders[1], Inches(4.1), Inches(1.0))
    for section in _sections(content):
        slide = deck.slides.add_slide(deck.slide_layouts[5])  # title only
        slide.shapes.title.text = section['title'] or title
        place(slide.shapes.title, Inches(0.3), Inches(1.0))
        item = next((items[r] for r in section['refs'] if r in items and not items[r].error), None)
        text_width = Inches(5.2) if item else Inches(12.3)
        if section['lines']:
            box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), text_width, Inches(5.5)).text_frame
            box.word_wrap = True
            for i, line in enumerate(section['lines'][:8]):
                paragraph = box.paragraphs[0] if i == 0 else box.add_paragraph()
                paragraph.text = f'• {line}'
                paragraph.font.size = Pt(16)
                paragraph.space_after = Pt(8)
        if item and item.png:
            left = Inches(5.9) if section['lines'] else Inches(1.5)
            slide.shapes.add_picture(io.BytesIO(item.png), left, Inches(1.5),
                                     width=Inches(7) if section['lines'] else Inches(10))
        elif item and item.result is not None:
            rows = item.result.rows[:10]
            columns = item.result.columns[:8]
            shape = slide.shapes.add_table(len(rows) + 1, len(columns), Inches(5.9), Inches(1.5), Inches(7), Inches(0.3))
            for i, name in enumerate(columns):
                shape.table.cell(0, i).text = name
            for r, row in enumerate(rows, 1):
                for i, value in enumerate(row[:8]):
                    shape.table.cell(r, i).text = _cell(value)
            for cell in (c for row in shape.table.rows for c in row.cells):
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.font.size = Pt(10)
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = 'Data source'
    box = slide.shapes.add_textbox(Inches(0.5), Inches(1.6), Inches(12.3), Inches(3)).text_frame
    box.word_wrap = True
    box.paragraphs[0].text = CREDIT
    box.paragraphs[0].font.size = Pt(16)
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


# ---- XLSX: a summary sheet, one sheet per table/plot result, and the SQL ----

def build_xlsx(title, content, items) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook()
    summary = book.active
    summary.title = 'Summary'
    summary.append([title])
    summary['A1'].font = Font(bold=True, size=14)
    summary.append([f'Generated by nembot on {datetime.now():%d %B %Y %H:%M} from AEMO MMS data.'])
    summary.append([CREDIT])
    summary.append([])
    for kind, value in blocks(content):
        if kind == 'text':
            for line in value.splitlines():
                summary.append([_plain(line)])
        else:
            summary.append([f'(see sheet {value})'])
    summary.column_dimensions['A'].width = 120
    sql_sheet = book.create_sheet('SQL')
    sql_sheet.append(['Result', 'SQL'])
    for item in items.values():
        sheet = book.create_sheet(item.id)
        if item.error or item.result is None:
            sheet.append([item.error or 'No data'])
        else:
            sheet.append(item.result.columns)
            for cell in sheet[1]:
                cell.font = Font(bold=True)
            for row in item.result.rows:
                sheet.append([_jsonable(v) if not isinstance(v, (int, float)) or isinstance(v, bool) else v for v in row])
            sheet.freeze_panes = 'A2'
        sql_sheet.append([item.id, item.sql])
    sql_sheet.column_dimensions['B'].width = 120
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ---- Notebooks: markdown and runnable cells that query through nembot.client ----

def _setup_code(url: str, database: str) -> str:
    return (f'import altair as alt\nfrom nembot.client import query\n\n'
            f'# Queries go to the running nembot app ({url}), or read {database} directly\n'
            f'# when the app is not running.\nNEMBOT = dict(url={url!r}, database={database!r})')


def _result_code(item: Item, suffix: str) -> str:
    code = f'df_{suffix} = query("""\n{item.sql}\n""", **NEMBOT)\n'
    if item.plot_spec:
        code += (f'chart_{suffix} = alt.Chart.from_dict({json.dumps(item.plot_spec, indent=1)}, validate=False)\n'
                 f'chart_{suffix}.data = df_{suffix}\nchart_{suffix}')
    else:
        code += f'df_{suffix}'
    return code


def build_jupyter(title, content, items, *, url: str, database: str) -> bytes:
    import nbformat
    from nbformat.v4 import new_code_cell, new_markdown_cell, new_notebook
    cells = [new_markdown_cell(f'# {title}\n\nGenerated by nembot on {datetime.now():%d %B %Y}.\n\n*{CREDIT}*'),
             new_code_cell(_setup_code(url, database))]
    for kind, value in blocks(content):
        if kind == 'text':
            cells.append(new_markdown_cell(value.replace('$', r'\$')))  # Jupyter reads $...$ as LaTeX
        elif value in items:
            cells.append(new_code_cell(_result_code(items[value], value.lower())))
    notebook = new_notebook(cells=cells, metadata={'kernelspec': {'name': 'python3', 'display_name': 'Python 3'}})
    return nbformat.writes(notebook).encode()


def build_marimo(title, content, items, *, url: str, database: str) -> bytes:
    def cell(body: str, uses: list[str], defines: list[str]) -> str:
        indented = '\n'.join('    ' + line if line else '' for line in body.splitlines())
        returns = f'    return ({", ".join(defines)},)' if defines else '    return'
        return f'@app.cell\ndef _({", ".join(uses)}):\n{indented}\n{returns}\n\n\n'

    def md(text: str) -> str:
        return cell(f'mo.md(r"""\n{text.replace(chr(34) * 3, chr(39) * 3)}\n""")', ['mo'], [])
    source = ('import marimo\n\n__generated_with = "nembot"\napp = marimo.App(width="medium")\n\n\n'
              + cell('import marimo as mo\n' + _setup_code(url, database), [], ['NEMBOT', 'alt', 'mo', 'query'])
              + md(f'# {title}\n\nGenerated by nembot on {datetime.now():%d %B %Y}.\n\n*{CREDIT}*'))
    for kind, value in blocks(content):
        if kind == 'text':
            source += md(value)
        elif value in items:
            item = items[value]
            source += cell(_result_code(item, value.lower()), ['NEMBOT', 'alt', 'query'] if item.plot_spec
                           else ['NEMBOT', 'query'], [])
    source += 'if __name__ == "__main__":\n    app.run()\n'
    return source.encode()


def export(db, directory: Path, fmt: str, title: str, content: str, results: dict, *,
           url: str, database: str) -> dict:
    """Build the file and return its metadata. Unreferenced results are omitted."""
    if fmt not in FORMATS:
        raise ValueError(f'format must be one of {", ".join(FORMATS)}')
    refs = [value for kind, value in blocks(content) if kind == 'result']
    if fmt == 'xlsx' and not refs:
        refs = list(results)  # a workbook of every result if none were placed
    missing = [r for r in refs if r not in results]
    items = prepare(db, refs, results, images=fmt in ('pdf', 'docx', 'pptx'))
    builders = {'html': lambda: build_html(title, content, items, static=False).encode(),
                'pdf': lambda: build_pdf(title, content, items), 'docx': lambda: build_docx(title, content, items),
                'pptx': lambda: build_pptx(title, content, items), 'xlsx': lambda: build_xlsx(title, content, items),
                'jupyter': lambda: build_jupyter(title, content, items, url=url, database=database),
                'marimo': lambda: build_marimo(title, content, items, url=url, database=database)}
    data = builders[fmt]()
    directory.mkdir(parents=True, exist_ok=True)
    name = f'{datetime.now():%Y%m%d-%H%M%S}-{slug(title)}{FORMATS[fmt]}'
    (directory / name).write_bytes(data)
    return {'file': name, 'format': fmt, 'title': title, 'bytes': len(data), 'missing': missing,
            'errors': {i.id: i.error for i in items.values() if i.error}}


def listing(directory: Path) -> list[dict]:
    if not directory.exists():
        return []
    files = []
    for path in sorted(directory.iterdir(), key=lambda p: p.stat().st_mtime, reverse=True):
        fmt = next((f for f, ext in FORMATS.items() if path.suffix == ext), None)
        if fmt and path.is_file():
            files.append({'file': path.name, 'format': fmt, 'bytes': path.stat().st_size,
                          'created': datetime.fromtimestamp(path.stat().st_mtime).isoformat(sep=' ', timespec='seconds')})
    return files
